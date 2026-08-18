import os
import re
import sys

def parse_md_table(md_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    headers = []
    rows = []
    in_table = False
    
    for line in lines:
        line_str = line.strip()
        if line_str.startswith('|') and line_str.endswith('|'):
            cells = [c.strip() for c in line_str.split('|')[1:-1]]
            # Check if separator row like |---|---|
            if all(set(c) <= set('-: ') for c in cells):
                in_table = True
                continue
            
            if not headers:
                headers = cells
            elif in_table:
                rows.append(cells)
        else:
            if in_table:
                in_table = False
                
    return headers, rows

def main():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("Installing openpyxl...")
        os.system(f"{sys.executable} -m pip install openpyxl")
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    subtitle_font = Font(name="Calibri", size=11, italic=True, color="595959")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    valid_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    invalid_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    incomplete_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    # 1. Summary Sheet
    ws_sum = wb.create_sheet(title="Summary")
    ws_sum.views.sheetView[0].showGridLines = True
    
    ws_sum.cell(row=1, column=1, value="HW06 — API Testing Summary Report").font = title_font
    ws_sum.cell(row=2, column=1, value="Student ID: 23127147 | System Under Test: EShop").font = subtitle_font
    
    # Table 1: Self Assessment
    ws_sum.cell(row=4, column=1, value="1. Self Assessment Table").font = Font(size=12, bold=True)
    sa_headers = ["No.", "Criteria", "Max Grade", "Self-Assessed Grade"]
    for col_idx, h in enumerate(sa_headers, 1):
        cell = ws_sum.cell(row=5, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sa_data = [
        [1, "API 1 — full pipeline (POST /api/login)", 30, 30],
        [2, "API 2 — full pipeline (POST/GET /api/cart)", 30, 28],
        [3, "API 3 — full pipeline (PUT /api/admin/orders/:id/status)", 30, 30],
        [4, "Agent Skills (AI-driven test generator design)", 10, 9],
        ["", "Total Grade", 100, 97]
    ]
    
    for row_offset, row_data in enumerate(sa_data, 6):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_sum.cell(row=row_offset, column=col_idx, value=val)
            cell.border = thin_border
            if row_offset == 10:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    # Table 2: Test Execution Metrics
    ws_sum.cell(row=12, column=1, value="2. Test Execution & Bug Summary").font = Font(size=12, bold=True)
    m_headers = ["Metric", "Value", "Description"]
    for col_idx, h in enumerate(m_headers, 1):
        cell = ws_sum.cell(row=13, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    m_data = [
        ["Total APIs Tested", 3, "Login (FR-02), Cart (FR-07), Admin Order Status (FR-18)"],
        ["AI-Generated Test Cases", 116, "Generated via Claude Sonnet 4.6"],
        ["Human Extended Test Cases", 15, "5 cases per API added manually"],
        ["Total Test Cases Designed", 131, "Total across 3 APIs"],
        ["Newman Requests Executed", 39, "Executed in Postman/Newman"],
        ["Newman Assertions Passed", 62, "100% assertions passed"],
        ["Newman Assertions Failed", 0, "Clean execution report"],
        ["Total Bugs Found", 14, "3 Critical, 6 High, 4 Medium, 1 Low"]
    ]
    
    for row_offset, row_data in enumerate(m_data, 14):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_sum.cell(row=row_offset, column=col_idx, value=val)
            cell.border = thin_border

    # Adjust summary column widths
    ws_sum.column_dimensions['A'].width = 30
    ws_sum.column_dimensions['B'].width = 45
    ws_sum.column_dimensions['C'].width = 15
    ws_sum.column_dimensions['D'].width = 25

    # 2. Test Case Sheets
    files_to_process = [
        ("API1 - Login", "test_cases/API1_Login_test_cases.md"),
        ("API2 - Cart", "test_cases/API2_Cart_test_cases.md"),
        ("API3 - Admin Order Status", "test_cases/API3_AdminOrderStatus_test_cases.md")
    ]

    for sheet_title, md_file in files_to_process:
        if not os.path.exists(md_file):
            continue
            
        headers, rows = parse_md_table(md_file)
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True
        
        # Header Row
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Data Rows
        for row_idx, r in enumerate(rows, 2):
            for col_idx, val in enumerate(r, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                
                # Highlight Audit Label
                if len(headers) >= col_idx and "Audit Label" in headers[col_idx-1]:
                    if "VALID" in val and "INVALID" not in val:
                        cell.fill = valid_fill
                    elif "INVALID" in val:
                        cell.fill = invalid_fill
                    elif "INCOMPLETE" in val:
                        cell.fill = incomplete_fill

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                lines = val_str.split('\n')
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    out_path = "test_cases/HW06_Test_Cases.xlsx"
    wb.save(out_path)
    print(f"Successfully generated Excel file at: {out_path}")

if __name__ == '__main__':
    main()
